"""Document loaders for PDF, TXT, and DOCX files producing LangChain Document objects."""

from pathlib import Path
from typing import List
from langchain_core.documents import Document
import pypdf
import docx

try:
    from backend.app.core.logging import get_logger
except ImportError:
    from app.core.logging import get_logger

logger = get_logger(__name__)


def load_pdf(file_path: Path, document_id: str, filename: str) -> List[Document]:
    """Extract text from a PDF file with page-level metadata using pypdf."""
    documents: List[Document] = []
    reader = pypdf.PdfReader(str(file_path))

    for page_idx, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        doc = Document(
            page_content=page_text,
            metadata={
                "document_id": document_id,
                "filename": filename,
                "file_type": "pdf",
                "page": page_idx + 1,
                "total_pages": len(reader.pages),
                "source": str(file_path),
            },
        )
        documents.append(doc)

    logger.info("Loaded %d pages from PDF file: %s", len(documents), filename)
    return documents


def load_txt(file_path: Path, document_id: str, filename: str) -> List[Document]:
    """Read a plain text document preserving document metadata."""
    try:
        text = file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        logger.warning("UTF-8 decoding failed for %s, falling back to latin-1", filename)
        text = file_path.read_text(encoding="latin-1", errors="replace")

    doc = Document(
        page_content=text,
        metadata={
            "document_id": document_id,
            "filename": filename,
            "file_type": "txt",
            "page": 1,
            "total_pages": 1,
            "source": str(file_path),
        },
    )
    logger.info("Loaded text file: %s (%d characters)", filename, len(text))
    return [doc]


def load_docx(file_path: Path, document_id: str, filename: str) -> List[Document]:
    """Extract paragraphs and table text from a DOCX file using python-docx."""
    doc_reader = docx.Document(str(file_path))
    content_parts: List[str] = []

    # Extract paragraphs
    for para in doc_reader.paragraphs:
        if para.text.strip():
            content_parts.append(para.text.strip())

    # Extract tables
    for table in doc_reader.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                content_parts.append(" | ".join(row_text))

    full_text = "\n\n".join(content_parts)
    doc = Document(
        page_content=full_text,
        metadata={
            "document_id": document_id,
            "filename": filename,
            "file_type": "docx",
            "page": 1,
            "total_pages": 1,
            "source": str(file_path),
        },
    )
    logger.info("Loaded DOCX file: %s (%d sections/paragraphs)", filename, len(content_parts))
    return [doc]


def load_document(file_path: Path, document_id: str, filename: str) -> List[Document]:
    """Dispatch file loading based on file extension."""
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        return load_pdf(file_path, document_id, filename)
    elif suffix == ".txt":
        return load_txt(file_path, document_id, filename)
    elif suffix == ".docx":
        return load_docx(file_path, document_id, filename)
    else:
        raise ValueError(f"Unsupported file format '{suffix}'. Supported: .pdf, .txt, .docx")
